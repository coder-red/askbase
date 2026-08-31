"""Unit tests for the document parser (HTML/PDF/XLSX/binary routing)."""

from __future__ import annotations

import io

import pytest

from basechatt.ingestion.parser import (
    _looks_binary,
    parse_bytes,
    parse_xlsx,
)


def _build_xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Region", "Revenue", "Growth"])
    ws.append(["Lagos", "1.2bn", "8%"])
    ws.append(["Abuja", "540m", "5%"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_xlsx_extracts_tables_and_sections():
    doc = parse_xlsx(_build_xlsx())
    assert doc.tables, "expected at least one table"
    table = doc.tables[0]
    assert table.title == "Sheet1"
    assert table.headers == ["Region", "Revenue", "Growth"]
    assert ["Lagos", "1.2bn", "8%"] in table.rows
    combined = " ".join(s.text for s in doc.sections)
    assert "Lagos" in combined and "1.2bn" in combined


def test_parse_bytes_routes_xlsx_mime():
    data = _build_xlsx()
    doc = parse_bytes(data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert doc.tables and doc.tables[0].headers == ["Region", "Revenue", "Growth"]


def test_parse_bytes_rejects_raw_zip_binary():
    doc = parse_bytes(b"PK\x03\x04" + b"\x00" * 100, "application/octet-stream")
    assert doc.tables == []
    assert doc.raw_text == ""


def test_parse_bytes_rejects_high_control_char_ratio():
    junk = bytes([1, 2, 3, 4]) * 50
    text = junk.decode("utf-8", errors="ignore")
    doc = parse_bytes(junk, "text/html")
    assert doc.raw_text == ""


def test_looks_binary_detects_zip_magic():
    assert _looks_binary(b"PK\x03\x04rest")
    assert _looks_binary(b"PK\x05\x06rest")
    assert not _looks_binary(b"<html>hi</html>")
